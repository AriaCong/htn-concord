# -*- coding: utf-8 -*-
"""Build the cv-llm-rulebench overview decks (ZH + EN).

Reuses the HTN-Concord slide functions read-only, so both decks share one look.
Usage: python build_deck.py ..      (writes the two .pptx files into docs/slides/)
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
HTN_SRC = os.path.join(HERE, "../../../../../_archive/htn-concord/docs/slides/src")
sys.path.insert(0, os.path.abspath(HTN_SRC))
sys.path.insert(0, HERE)

import build as hb            # HTN-Concord slide functions (unchanged)
from deck_lib import new_deck
import content_rb as C

def cover_rb(prs, L):
    """HTN-Concord cover, with tag pills sized for CJK text (the shared estimate wraps them)."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    s = hb.base_slide(prs)
    hb.rect(s, 0, 0, 0.30, hb.H, fill=hb.ACCENT, line=None)
    hb.rect(s, 0, hb.H - 1.55, 0.30, 1.55, fill=hb.AMBER, line=None)
    tf = hb.textbox(s, 1.20, 1.55, 11.0, 0.3)
    hb.para(tf, L["cover_kicker"], size=12, color=hb.ACCENT, bold=True, first=True, space_after=0)
    tf = hb.textbox(s, 1.20, 2.00, 11.0, 1.9)
    for i, ln in enumerate(L["cover_title"].split("\n")):
        hb.para(tf, ln, size=40, color=hb.INK, bold=True, first=(i == 0), space_after=2, line=1.12)
    hb.hline(s, 1.20, 4.06, 2.2, hb.AMBER, 3.0)
    tf = hb.textbox(s, 1.20, 4.32, 10.4, 0.9)
    hb.para(tf, L["cover_sub"], size=16, color=hb.SLATE, first=True, space_after=0, line=1.35)
    x = 1.20
    for t in L["cover_tags"]:
        cjk = sum(1 for ch in t if ord(ch) > 0x2E80)
        w = cjk * 0.19 + (len(t) - cjk) * 0.10 + 0.50
        b = hb.rect(s, x, 5.42, w, 0.36, fill=hb.CARD2, line=None,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        tfb = b.text_frame
        tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
        tfb.margin_left = tfb.margin_right = 0
        tfb.word_wrap = False
        hb.para(tfb, t, size=11, color=hb.ACCENT_D, bold=True, first=True, space_after=0,
                align=PP_ALIGN.CENTER)
        x += w + 0.16
    tf = hb.textbox(s, 1.20, 6.36, 10.4, 0.5)
    hb.para(tf, "   ·   ".join(L["cover_meta"]), size=12, color=hb.MUTED, first=True, space_after=0)


def s14_results(prs, L, n):
    """Pilot results: same look as the HTN-Concord grid slide, wider failure column."""
    s = hb.base_slide(prs, L["s14_t"], L["s14_k"], n, hb.FOOT)
    rows = [(b, fail, acc, verdict) for b, acc, fail, verdict in L["s14_rows"]]
    end = hb.grid(s, L, hb.ML, hb.BODY_TOP + 0.06, hb.CW, [0.85, 5.2, 2.4, 1.9],
                  rows, size=12, badge=True, first_bold=False,
                  colcolor={1: hb.INK, 3: hb.ACCENT_D})
    tf = hb.textbox(s, hb.ML, end + 0.26, hb.CW, 0.6)
    hb.para(tf, L["s14_note"], size=12, color=hb.SLATE, first=True, space_after=0, line=1.38)
    hb.check("s14", end + 0.26 + hb.est_h(L["s14_note"], hb.CW, 12, L["lang"], 1.38, 0, 0),
             hb.BODY_BOT + 0.2)


SLIDES = [hb.s2_quad, hb.s3_constraints, hb.s4_gap, hb.s5_positioning, hb.s6_engine,
          hb.s7_subtraction, hb.s8_arch, hb.s9_scope, hb.s10_data, hb.s11_tasks,
          hb.s12_ladder, hb.s13_metrics, s14_results, hb.s15_conclusions,
          hb.s16_progress, hb.s17_roadmap, hb.s18_risks, hb.s19_contrib]


def build(L, outpath):
    hb.FOOT = L["foot"]
    prs = new_deck()
    cover_rb(prs, L)
    for i, fn in enumerate(SLIDES, start=2):
        fn(prs, L, i)
    prs.save(outpath)
    return len(prs.slides)


if __name__ == "__main__":
    outdir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..")
    for L in (C.ZH, C.EN):
        hb.WARN[:] = []
        p = os.path.join(outdir, L["file"])
        n = build(L, p)
        print("%-32s %2d slides  %6.1f KB" % (L["file"], n, os.path.getsize(p) / 1024))
        for w in hb.WARN:
            print("    OVERFLOW  " + w)
        if not hb.WARN:
            print("    layout check: no overflow")
