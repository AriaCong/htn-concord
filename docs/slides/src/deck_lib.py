"""Layout primitives for the HTN-Concord supervisor deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette
INK      = RGBColor(0x18, 0x25, 0x33)
SLATE    = RGBColor(0x47, 0x5A, 0x6B)
MUTED    = RGBColor(0x74, 0x86, 0x94)
ACCENT   = RGBColor(0x0E, 0x74, 0x90)   # teal
ACCENT_D = RGBColor(0x09, 0x4F, 0x63)
AMBER    = RGBColor(0xB4, 0x53, 0x09)
GREEN    = RGBColor(0x15, 0x7F, 0x5C)
CARD     = RGBColor(0xF2, 0xF6, 0xF8)
CARD2    = RGBColor(0xE7, 0xEF, 0xF3)
LINE     = RGBColor(0xD3, 0xDF, 0xE6)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

W, H = 13.333, 7.5
ML, MR = 0.78, 0.78
CW = W - ML - MR          # content width
BODY_TOP = 1.62
BODY_BOT = 6.82

LATIN = "Arial"
EA    = "Microsoft YaHei"


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def _set_fonts(run):
    """Force the east-asian typeface too, so CJK does not fall back oddly."""
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", LATIN), ("a:ea", EA), ("a:cs", LATIN)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=14, color=INK, bold=False, space_before=0, space_after=6,
         align=None, line=1.22, first=False, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    _set_fonts(r)
    return p


def rich(tf, chunks, size=14, space_before=0, space_after=6, align=None,
         line=1.22, first=False):
    """chunks: list of (text, color, bold)"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    if align is not None:
        p.alignment = align
    for text, color, bold in chunks:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        _set_fonts(r)
    return p


def rect(slide, x, y, w, h, fill=CARD, line=None, lw=1.0, radius=None,
         shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.16)
    tf.margin_top = tf.margin_bottom = Inches(0.11)
    tf.paragraphs[0].text = ""
    return sh


def hline(slide, x, y, w, color=LINE, lw=1.0):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(lw)
    return ln


def arrow(slide, x, y, w, h, color=ACCENT, shape=MSO_SHAPE.RIGHT_ARROW):
    a = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


# ---------------------------------------------------------------- chrome
def base_slide(prs, title=None, kicker=None, n=None, footer=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    if kicker:
        tf = textbox(s, ML, 0.42, CW, 0.26)
        para(tf, kicker.upper(), size=10.5, color=ACCENT, bold=True, first=True,
             space_after=0)
    if title:
        tf = textbox(s, ML, 0.70, CW, 0.62)
        para(tf, title, size=27, color=INK, bold=True, first=True, space_after=0,
             line=1.05)
        hline(s, ML, 1.42, 1.5, ACCENT, 2.4)
        hline(s, ML + 1.5, 1.42, CW - 1.5, LINE, 1.0)
    if n is not None:
        tf = textbox(s, W - MR - 1.0, H - 0.52, 1.0, 0.26, align=PP_ALIGN.RIGHT)
        para(tf, str(n), size=9.5, color=MUTED, first=True, space_after=0)
        tf2 = textbox(s, ML, H - 0.52, 8.0, 0.26)
        para(tf2, footer, size=9.5, color=MUTED, first=True, space_after=0)
    return s
