# -*- coding: utf-8 -*-
import math, sys, os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from deck_lib import *
import content as C

WARN = []

# ---------------------------------------------------------------- metrics
def _cpi(size, lang):
    """Approximate characters per inch for a given point size."""
    base = 7.0 if lang == "zh" else 13.2      # at 11pt
    return base * 11.0 / size

def est_lines(text, width_in, size, lang, pad_in=0.34):
    usable = max(width_in - pad_in, 0.4)
    cpi = _cpi(size, lang)
    n = 0
    for seg in text.split("\n"):
        n += max(1, math.ceil(len(seg) / max(usable * cpi, 1)))
    return n

def est_h(text, width_in, size, lang, line=1.24, pad_in=0.34, vpad=0.24):
    return est_lines(text, width_in, size, lang, pad_in) * size * line / 72.0 + vpad

def check(name, need, have):
    if need > have + 0.02:
        WARN.append("%-26s needs %.2f\" has %.2f\"" % (name, need, have))

# ---------------------------------------------------------------- slides
def cover(prs, L):
    s = base_slide(prs)
    rect(s, 0, 0, 0.30, H, fill=ACCENT, line=None)
    rect(s, 0, H - 1.55, 0.30, 1.55, fill=AMBER, line=None)
    tf = textbox(s, 1.20, 1.55, 11.0, 0.3)
    para(tf, L["cover_kicker"], size=12, color=ACCENT, bold=True, first=True, space_after=0)
    tf = textbox(s, 1.20, 2.00, 11.0, 1.9)
    for i, ln in enumerate(L["cover_title"].split("\n")):
        para(tf, ln, size=40, color=INK, bold=True, first=(i == 0), space_after=2, line=1.12)
    hline(s, 1.20, 4.06, 2.2, AMBER, 3.0)
    tf = textbox(s, 1.20, 4.32, 10.4, 0.9)
    para(tf, L["cover_sub"], size=16, color=SLATE, first=True, space_after=0, line=1.35)
    x = 1.20
    for t in L["cover_tags"]:
        w = est_lines(t, 9, 11, L["lang"], 0) and (len(t) / _cpi(11, L["lang"]) + 0.44)
        b = rect(s, x, 5.42, w, 0.36, fill=CARD2, line=None,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        tfb = b.text_frame
        tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tfb, t, size=11, color=ACCENT_D, bold=True, first=True, space_after=0,
             align=PP_ALIGN.CENTER)
        x += w + 0.16
    tf = textbox(s, 1.20, 6.36, 10.4, 0.5)
    para(tf, "   ·   ".join(L["cover_meta"]), size=12, color=MUTED, first=True, space_after=0)


def s2_quad(prs, L, n):
    s = base_slide(prs, L["s2_t"], L["s2_k"], n, FOOT)
    gap = 0.30
    cw = (CW - gap) / 2
    ch = (BODY_BOT - BODY_TOP - gap) / 2
    for i, (head, body) in enumerate(L["s2"]):
        x = ML + (i % 2) * (cw + gap)
        y = BODY_TOP + (i // 2) * (ch + gap)
        b = rect(s, x, y, cw, ch, fill=CARD, line=None)
        rect(s, x, y, 0.055, ch, fill=ACCENT if i < 2 else AMBER, line=None)
        tf = textbox(s, x + 0.32, y + 0.28, cw - 0.62, ch - 0.5)
        para(tf, head, size=17, color=ACCENT_D if i < 2 else AMBER, bold=True,
             first=True, space_after=9)
        para(tf, body, size=13.5, color=INK, space_after=0, line=1.42)
        check("s2-%d" % i, est_h(body, cw - 0.62, 13.5, L["lang"], 1.42, 0) + 0.55, ch - 0.28)


def s3_constraints(prs, L, n):
    s = base_slide(prs, L["s3_t"], L["s3_k"], n, FOOT)
    tf = textbox(s, ML, BODY_TOP, CW, 0.36)
    para(tf, L["s3_lead"], size=13.5, color=SLATE, first=True, space_after=0)
    gap = 0.24
    cw = (CW - 3 * gap) / 4
    y = BODY_TOP + 0.62
    ch = 3.05
    for i, (head, fact, miss) in enumerate(L["s3_cols"]):
        x = ML + i * (cw + gap)
        rect(s, x, y, cw, ch, fill=CARD, line=None)
        rect(s, x, y, cw, 0.05, fill=ACCENT, line=None)
        tf = textbox(s, x + 0.26, y + 0.30, cw - 0.52, 0.4)
        para(tf, head, size=14.5, color=INK, bold=True, first=True, space_after=0)
        tf = textbox(s, x + 0.26, y + 0.86, cw - 0.52, 1.2)
        for j, ln in enumerate(fact.split("\n")):
            para(tf, ln, size=13, color=ACCENT_D, bold=True, first=(j == 0),
                 space_after=2, line=1.32)
        rect(s, x + 0.26, y + ch - 1.02, cw - 0.52, 0.02, fill=LINE, line=None)
        tf = textbox(s, x + 0.26, y + ch - 0.86, cw - 0.52, 0.72)
        para(tf, miss, size=11.5, color=AMBER, bold=True, first=True, space_after=0, line=1.3)
        check("s3-%d" % i, est_h(miss, cw - 0.52, 11.5, L["lang"], 1.3, 0), 0.78)
    tf = textbox(s, ML, y + ch + 0.30, CW, 0.6)
    para(tf, L["s3_foot"], size=12.5, color=MUTED, first=True, space_after=0, line=1.35)


def s4_gap(prs, L, n):
    s = base_slide(prs, L["s4_t"], L["s4_k"], n, FOOT)
    lw = CW * 0.585
    y = BODY_TOP + 0.10
    for head, sub in L["s4_bullets"]:
        indent = head.startswith("  ")
        h1 = est_h(head, lw - (0.34 if indent else 0.0), 14.5, L["lang"], 1.3, 0, 0)
        tf = textbox(s, ML + (0.34 if indent else 0.0), y,
                     lw - (0.34 if indent else 0.0), h1 + 0.1)
        para(tf, head.strip() if indent else head, size=14.5,
             color=INK, bold=not indent, first=True, space_after=0, line=1.3)
        if indent:
            rect(s, ML + 0.10, y + 0.09, 0.06, 0.06, fill=ACCENT, line=None,
                 shape=MSO_SHAPE.OVAL)
        y += h1 + 0.04
        if sub:
            h2 = est_h(sub, lw - 0.34, 12.5, L["lang"], 1.3, 0, 0)
            tf = textbox(s, ML + 0.34, y, lw - 0.34, h2 + 0.1)
            para(tf, sub, size=12.5, color=MUTED, first=True, space_after=0, line=1.3)
            y += h2 + 0.14
        else:
            y += 0.06
    check("s4-left", y, BODY_BOT)

    bx = ML + lw + 0.42
    bw = CW - lw - 0.42
    bh = est_h(L["s4_box"], bw, 15, L["lang"], 1.45, 0.7) + 0.85
    b = rect(s, bx, BODY_TOP + 0.55, bw, bh, fill=ACCENT_D, line=None)
    tf = textbox(s, bx + 0.36, BODY_TOP + 0.55 + 0.34, bw - 0.72, 0.32)
    para(tf, L["s4_box_t"], size=11.5, color=RGBColor(0x8E, 0xD6, 0xE4), bold=True,
         first=True, space_after=0)
    tf = textbox(s, bx + 0.36, BODY_TOP + 0.55 + 0.82, bw - 0.72, bh - 1.1)
    para(tf, L["s4_box"], size=15, color=WHITE, bold=True, first=True, space_after=0, line=1.45)


def grid(s, L, x, y, w, cols, rows, header=None, size=11.5, head_size=11,
         first_bold=True, badge=False, zebra=True, maxh=None, colcolor=None):
    """cols = list of relative widths. rows = list of tuples of strings."""
    tot = sum(cols)
    widths = [w * c / tot for c in cols]
    cy = y
    if header:
        hh = 0.40
        rect(s, x, cy, w, hh, fill=INK, line=None)
        cx = x
        for i, cell in enumerate(header):
            tf = textbox(s, cx + 0.16, cy + 0.10, widths[i] - 0.32, hh - 0.16)
            para(tf, cell, size=head_size, color=WHITE, bold=True, first=True, space_after=0)
            cx += widths[i]
        cy += hh
    for ri, row in enumerate(rows):
        hs = []
        for i, cell in enumerate(row):
            sz = size + (0.5 if (i == 0 and first_bold) else 0)
            hs.append(est_h(cell, widths[i], sz, L["lang"], 1.28, 0.32, 0.26))
        rh = max(hs + [0.44])
        if maxh:
            rh = min(rh, maxh)
        if zebra and ri % 2 == 1:
            rect(s, x, cy, w, rh, fill=CARD, line=None)
        else:
            rect(s, x, cy, w, rh, fill=WHITE, line=None)
        hline(s, x, cy, w, LINE, 0.75)
        cx = x
        for i, cell in enumerate(row):
            if i == 0 and badge:
                bw = min(widths[0] - 0.3, 0.86)
                bb = rect(s, cx + 0.14, cy + (rh - 0.32) / 2, bw, 0.32, fill=ACCENT,
                          line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
                tfb = bb.text_frame
                tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
                tfb.margin_left = tfb.margin_right = 0
                para(tfb, cell, size=11, color=WHITE, bold=True, first=True,
                     space_after=0, align=PP_ALIGN.CENTER)
            else:
                col = INK if (i == 0 and first_bold) else SLATE
                if colcolor and i in colcolor:
                    col = colcolor[i]
                tf = textbox(s, cx + 0.16, cy + 0.13, widths[i] - 0.32, rh - 0.2)
                para(tf, cell, size=size + (0.5 if (i == 0 and first_bold) else 0),
                     color=col, bold=(i == 0 and first_bold), first=True,
                     space_after=0, line=1.28)
            cx += widths[i]
        cy += rh
    hline(s, x, cy, w, LINE, 0.75)
    return cy


def s5_positioning(prs, L, n):
    s = base_slide(prs, L["s5_t"], L["s5_k"], n, FOOT)
    end = grid(s, L, ML, BODY_TOP + 0.05, CW, [2.5, 4.0, 3.6],
               L["s5_rows"], size=12.5,
               colcolor={2: ACCENT_D})
    check("s5", end, BODY_BOT)


def s6_engine(prs, L, n):
    s = base_slide(prs, L["s6_t"], L["s6_k"], n, FOOT)
    y = BODY_TOP + 0.05
    inw, arw, outw = 3.35, 0.62, CW - 3.35 - 0.62 - 0.50
    box_h = max(1.62, 0.62 + sum(est_h(i, outw - 0.9, 13, L["lang"], 1.3, 0, 0) + 0.10
                                 for i in L["s6_out"]) + 0.26)
    b = rect(s, ML, y, inw, box_h, fill=CARD2, line=None)
    tf = textbox(s, ML + 0.26, y + 0.24, inw - 0.52, 0.3)
    para(tf, L["s6_in_t"], size=11, color=ACCENT_D, bold=True, first=True, space_after=0)
    tf = textbox(s, ML + 0.26, y + 0.66, inw - 0.52, box_h - 0.9)
    for i, ln in enumerate(L["s6_in"].split("\n")):
        para(tf, ln, size=13, color=INK, first=(i == 0), space_after=3, line=1.32)
    arrow(s, ML + inw + 0.24, y + box_h / 2 - 0.16, arw - 0.06, 0.32, ACCENT)
    ox = ML + inw + arw + 0.50
    rect(s, ox, y, outw, box_h, fill=WHITE, line=ACCENT, lw=1.6)
    tf = textbox(s, ox + 0.28, y + 0.22, outw - 0.56, 0.3)
    para(tf, L["s6_out_t"], size=11, color=ACCENT_D, bold=True, first=True, space_after=0)
    oy = y + 0.62
    for item in L["s6_out"]:
        hh = est_h(item, outw - 0.9, 13, L["lang"], 1.3, 0, 0)
        rect(s, ox + 0.30, oy + 0.09, 0.055, 0.055, fill=AMBER, line=None, shape=MSO_SHAPE.OVAL)
        tf = textbox(s, ox + 0.52, oy, outw - 0.9, hh + 0.1)
        para(tf, item, size=13, color=INK, first=True, space_after=0, line=1.3)
        oy += hh + 0.10
    check("s6-out", oy - y, box_h)

    y2 = y + box_h + 0.42
    gap = 0.26
    cw = (CW - 2 * gap) / 3
    ch = BODY_BOT - y2
    for i, (head, body) in enumerate(L["s6_rules"]):
        x = ML + i * (cw + gap)
        rect(s, x, y2, cw, ch, fill=CARD, line=None)
        rect(s, x, y2, cw, 0.045, fill=[ACCENT, AMBER, GREEN][i], line=None)
        tf = textbox(s, x + 0.26, y2 + 0.30, cw - 0.52, 0.44)
        para(tf, head, size=14, color=INK, bold=True, first=True, space_after=0, line=1.2)
        hh = est_h(head, cw - 0.52, 14, L["lang"], 1.2, 0, 0)
        tf = textbox(s, x + 0.26, y2 + 0.34 + hh + 0.16, cw - 0.52, ch - hh - 0.7)
        para(tf, body, size=12, color=SLATE, first=True, space_after=0, line=1.38)
        check("s6r-%d" % i, hh + est_h(body, cw - 0.52, 12, L["lang"], 1.38, 0, 0) + 0.6, ch)


def s7_subtraction(prs, L, n):
    s = base_slide(prs, L["s7_t"], L["s7_k"], n, FOOT)
    y = BODY_TOP + 0.02
    gap = 0.34
    cw = (CW - gap) / 2
    ch = 2.05
    for i, (t, body, note, col) in enumerate([
            (L["s7_a_t"], L["s7_a"], L["s7_a_note"], ACCENT),
            (L["s7_b_t"], L["s7_b"], L["s7_b_note"], AMBER)]):
        x = ML + i * (cw + gap)
        rect(s, x, y, cw, ch, fill=WHITE, line=col, lw=1.6)
        rect(s, x, y, cw, 0.34, fill=col, line=None)
        tf = textbox(s, x + 0.24, y + 0.06, cw - 0.48, 0.24)
        para(tf, t, size=11.5, color=WHITE, bold=True, first=True, space_after=0)
        tf = textbox(s, x + 0.26, y + 0.52, cw - 0.52, ch - 1.05)
        for j, ln in enumerate(body.split("\n")):
            para(tf, ln, size=12, color=INK, first=(j == 0), space_after=2, line=1.36)
        hline(s, x + 0.26, y + ch - 0.44, cw - 0.52, LINE, 0.75)
        tf = textbox(s, x + 0.26, y + ch - 0.36, cw - 0.52, 0.28)
        para(tf, note, size=11.5, color=col, bold=True, first=True, space_after=0)

    yb = y + ch + 0.34
    rect(s, ML, yb, CW, 0.78, fill=INK, line=None)
    tf = textbox(s, ML, yb + 0.20, CW, 0.42, align=PP_ALIGN.CENTER)
    para(tf, L["s7_eq"], size=17, color=WHITE, bold=True, first=True, space_after=0,
         align=PP_ALIGN.CENTER)

    yy = yb + 0.78 + 0.32
    for item in L["s7_why"]:
        hh = est_h(item, CW - 0.42, 13, L["lang"], 1.34, 0, 0)
        rect(s, ML + 0.02, yy + 0.10, 0.06, 0.06, fill=ACCENT, line=None, shape=MSO_SHAPE.OVAL)
        tf = textbox(s, ML + 0.30, yy, CW - 0.30, hh + 0.1)
        para(tf, item, size=13, color=SLATE, first=True, space_after=0, line=1.34)
        yy += hh + 0.13
    check("s7", yy, BODY_BOT + 0.12)


def s8_arch(prs, L, n):
    s = base_slide(prs, L["s8_t"], L["s8_k"], n, FOOT)
    y = BODY_TOP - 0.02
    # sources
    gap = 0.49
    sw = (CW - 2 * gap) / 3
    for i, txt in enumerate(L["s8_sources"]):
        x = ML + i * (sw + gap)
        rect(s, x, y, sw, 0.66, fill=CARD2, line=None)
        tf = textbox(s, x + 0.18, y + 0.11, sw - 0.36, 0.48, align=PP_ALIGN.CENTER)
        for j, ln in enumerate(txt.split("\n")):
            para(tf, ln, size=(12 if j == 0 else 10.5),
                 color=(INK if j == 0 else MUTED), bold=(j == 0), first=(j == 0),
                 space_after=0, line=1.2, align=PP_ALIGN.CENTER)
    arrow(s, W / 2 - 0.16, y + 0.72, 0.32, 0.30, ACCENT, MSO_SHAPE.DOWN_ARROW)
    # profile
    y2 = y + 1.08
    rect(s, ML, y2, CW, 0.72, fill=ACCENT_D, line=None)
    tf = textbox(s, ML, y2 + 0.11, CW, 0.52, align=PP_ALIGN.CENTER)
    lines = L["s8_profile"].split("\n")
    para(tf, lines[0], size=14, color=WHITE, bold=True, first=True, space_after=1,
         align=PP_ALIGN.CENTER)
    para(tf, lines[1], size=10.5, color=RGBColor(0xA8, 0xD8, 0xE4), space_after=0,
         align=PP_ALIGN.CENTER, line=1.2)
    # two branches
    colw = (CW - 0.57) / 2
    lx, rx = ML, ML + colw + 0.57
    y3 = y2 + 0.72 + 0.34
    arrow(s, lx + colw / 2 - 0.14, y2 + 0.76, 0.28, 0.26, ACCENT, MSO_SHAPE.DOWN_ARROW)
    arrow(s, rx + colw / 2 - 0.14, y2 + 0.76, 0.28, 0.26, AMBER, MSO_SHAPE.DOWN_ARROW)

    def branch(x, top, mid, bot, col, dim):
        rect(s, x, top, colw, 0.62, fill=WHITE, line=col, lw=1.6)
        tf = textbox(s, x + 0.2, top + 0.09, colw - 0.4, 0.46, align=PP_ALIGN.CENTER)
        for j, ln in enumerate(mid.split("\n")):
            para(tf, ln, size=(13.5 if j == 0 else 10.5), color=(INK if j == 0 else MUTED),
                 bold=(j == 0), first=(j == 0), space_after=0, align=PP_ALIGN.CENTER, line=1.2)
        arrow(s, x + colw / 2 - 0.14, top + 0.66, 0.28, 0.24, col, MSO_SHAPE.DOWN_ARROW)
        b2 = top + 0.94
        rect(s, x, b2, colw, 0.66, fill=dim, line=None)
        tf = textbox(s, x + 0.2, b2 + 0.11, colw - 0.4, 0.5, align=PP_ALIGN.CENTER)
        for j, ln in enumerate(bot.split("\n")):
            para(tf, ln, size=(12.5 if j == 0 else 10.5), color=(INK if j == 0 else col),
                 bold=True, first=(j == 0), space_after=0, align=PP_ALIGN.CENTER, line=1.2)
        return b2 + 0.66

    e1 = branch(lx, y3, L["s8_engine"], L["s8_engine_out"], ACCENT, CARD)
    e2 = branch(rx, y3, L["s8_render"], L["s8_llm"], AMBER, CARD)
    ye = max(e1, e2) + 0.30
    arrow(s, lx + colw / 2 - 0.14, ye - 0.28, 0.28, 0.24, SLATE, MSO_SHAPE.DOWN_ARROW)
    arrow(s, rx + colw / 2 - 0.14, ye - 0.28, 0.28, 0.24, SLATE, MSO_SHAPE.DOWN_ARROW)
    rect(s, ML, ye, CW, 0.56, fill=INK, line=None)
    tf = textbox(s, ML, ye + 0.14, CW, 0.34, align=PP_ALIGN.CENTER)
    para(tf, L["s8_eval"], size=13.5, color=WHITE, bold=True, first=True, space_after=0,
         align=PP_ALIGN.CENTER)
    tf = textbox(s, ML, ye + 0.70, CW, 0.5)
    para(tf, L["s8_gate"], size=11, color=MUTED, first=True, space_after=0, line=1.3)
    check("s8", ye + 0.70 + est_h(L["s8_gate"], CW, 11, L["lang"], 1.3, 0, 0), BODY_BOT + 0.15)


def s9_scope(prs, L, n):
    s = base_slide(prs, L["s9_t"], L["s9_k"], n, FOOT)
    y = BODY_TOP + 0.02
    gap = 0.34
    cw = (CW - gap) / 2
    ch = 3.55
    for i, (t, items, col) in enumerate([(L["s9_in_t"], L["s9_in"], GREEN),
                                         (L["s9_out_t"], L["s9_out"], AMBER)]):
        x = ML + i * (cw + gap)
        rect(s, x, y, cw, ch, fill=CARD, line=None)
        rect(s, x, y, cw, 0.42, fill=col, line=None)
        tf = textbox(s, x + 0.26, y + 0.10, cw - 0.52, 0.3)
        para(tf, t, size=13, color=WHITE, bold=True, first=True, space_after=0)
        yy = y + 0.66
        for it in items:
            hh = est_h(it, cw - 0.86, 13, L["lang"], 1.3, 0, 0)
            tf = textbox(s, x + 0.30, yy, 0.24, 0.3)
            para(tf, "✓" if i == 0 else "✕", size=12, color=col, bold=True, first=True,
                 space_after=0)
            tf = textbox(s, x + 0.60, yy, cw - 0.9, hh + 0.1)
            para(tf, it, size=13, color=INK, first=True, space_after=0, line=1.3)
            yy += hh + 0.18
        check("s9-%d" % i, yy - y, ch)
    tf = textbox(s, ML, y + ch + 0.32, CW, 0.7)
    para(tf, L["s9_note"], size=12.5, color=SLATE, first=True, space_after=0, line=1.38)


def s10_data(prs, L, n):
    s = base_slide(prs, L["s10_t"], L["s10_k"], n, FOOT)
    y = BODY_TOP + 0.02
    gap = 0.34
    cw = (CW - gap) / 2
    ch = 3.55
    for i, (t, items, status) in enumerate(L["s10_cards"]):
        x = ML + i * (cw + gap)
        col = ACCENT if i == 0 else AMBER
        rect(s, x, y, cw, ch, fill=WHITE, line=LINE, lw=1.2)
        rect(s, x, y, cw, 0.055, fill=col, line=None)
        tf = textbox(s, x + 0.28, y + 0.30, cw - 0.56, 0.36)
        para(tf, t, size=17, color=INK, bold=True, first=True, space_after=0)
        pw = len(status) / _cpi(10.5, L["lang"]) + 0.4
        pb = rect(s, x + cw - 0.28 - pw, y + 0.32, pw, 0.30, fill=CARD2, line=None,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        tfb = pb.text_frame
        tfb.vertical_anchor = MSO_ANCHOR.MIDDLE
        tfb.margin_left = tfb.margin_right = 0
        para(tfb, status, size=10.5, color=ACCENT_D, bold=True, first=True, space_after=0,
             align=PP_ALIGN.CENTER)
        yy = y + 0.86
        for it in items:
            hh = est_h(it, cw - 0.86, 13, L["lang"], 1.32, 0, 0)
            rect(s, x + 0.32, yy + 0.09, 0.055, 0.055, fill=col, line=None, shape=MSO_SHAPE.OVAL)
            tf = textbox(s, x + 0.56, yy, cw - 0.86, hh + 0.1)
            para(tf, it, size=13, color=SLATE, first=True, space_after=0, line=1.32)
            yy += hh + 0.16
        check("s10-%d" % i, yy - y, ch)
    tf = textbox(s, ML, y + ch + 0.32, CW, 0.6)
    para(tf, L["s10_note"], size=12.5, color=SLATE, first=True, space_after=0, line=1.38)


def s11_tasks(prs, L, n):
    s = base_slide(prs, L["s11_t"], L["s11_k"], n, FOOT)
    end = grid(s, L, ML, BODY_TOP + 0.12, CW, [0.7, 2.2, 3.5, 3.6, 1.3],
               L["s11_rows"], size=12.5, badge=True, first_bold=False,
               colcolor={1: INK, 4: AMBER})
    check("s11", end, BODY_BOT)


def s12_ladder(prs, L, n):
    s = base_slide(prs, L["s12_t"], L["s12_k"], n, FOOT)
    end = grid(s, L, ML, BODY_TOP + 0.06, CW, [0.75, 3.9, 2.5, 3.0],
               L["s12_rows"], size=12.5, badge=True, first_bold=False,
               colcolor={1: INK, 3: ACCENT_D})
    tf = textbox(s, ML, end + 0.28, CW, 0.6)
    para(tf, L["s12_note"], size=12.5, color=SLATE, first=True, space_after=0, line=1.38)
    check("s12", end + 0.28 + est_h(L["s12_note"], CW, 12.5, L["lang"], 1.38, 0, 0), BODY_BOT + 0.2)


def s13_metrics(prs, L, n):
    s = base_slide(prs, L["s13_t"], L["s13_k"], n, FOOT)
    y = BODY_TOP + 0.02
    lw = CW * 0.44
    yy = y
    for t, items, col in [(L["s13_m1_t"], L["s13_m1"], ACCENT),
                          (L["s13_m2_t"], L["s13_m2"], AMBER),
                          (L["s13_m3_t"], [" · ".join(L["s13_m3"])], SLATE)]:
        bh = 0.46 + sum(est_h(i, lw - 0.9, 12.5, L["lang"], 1.28, 0, 0) + 0.10
                        for i in items) + 0.22
        rect(s, ML, yy, lw, bh, fill=CARD, line=None)
        rect(s, ML, yy, 0.05, bh, fill=col, line=None)
        tf = textbox(s, ML + 0.28, yy + 0.16, lw - 0.5, 0.28)
        para(tf, t, size=11.5, color=col, bold=True, first=True, space_after=0)
        iy = yy + 0.52
        for it in items:
            hh = est_h(it, lw - 0.9, 12.5, L["lang"], 1.28, 0, 0)
            tf = textbox(s, ML + 0.28, iy, lw - 0.56, hh + 0.1)
            para(tf, "· " + it, size=12.5, color=INK, first=True, space_after=0, line=1.28)
            iy += hh + 0.10
        yy += bh + 0.22
    check("s13-left", yy, BODY_BOT + 0.1)

    rx = ML + lw + 0.44
    rw = CW - lw - 0.44
    ry = y
    for i, (head, body) in enumerate(L["s13_rules"]):
        hh1 = est_h(head, rw - 0.5, 14, L["lang"], 1.22, 0, 0)
        hh2 = est_h(body, rw - 0.5, 12, L["lang"], 1.36, 0, 0)
        bh = hh1 + hh2 + 0.62
        rect(s, rx, ry, rw, bh, fill=WHITE, line=LINE, lw=1.2)
        tf = textbox(s, rx + 0.28, ry + 0.20, 0.3, 0.3)
        para(tf, str(i + 1), size=13, color=ACCENT, bold=True, first=True, space_after=0)
        tf = textbox(s, rx + 0.60, ry + 0.20, rw - 0.9, hh1 + 0.1)
        para(tf, head, size=14, color=INK, bold=True, first=True, space_after=0, line=1.22)
        tf = textbox(s, rx + 0.60, ry + 0.24 + hh1 + 0.08, rw - 0.9, hh2 + 0.1)
        para(tf, body, size=12, color=SLATE, first=True, space_after=0, line=1.36)
        ry += bh + 0.20
    check("s13-right", ry, BODY_BOT + 0.1)


def s14_expected(prs, L, n):
    s = base_slide(prs, L["s14_t"], L["s14_k"], n, FOOT)
    end = grid(s, L, ML, BODY_TOP + 0.06, CW, [0.6, 3.7, 3.4, 1.9],
               L["s14_rows"], size=12.5, badge=True, first_bold=False,
               colcolor={1: INK, 3: ACCENT_D})
    tf = textbox(s, ML, end + 0.28, CW, 0.6)
    para(tf, L["s14_note"], size=12.5, color=SLATE, first=True, space_after=0, line=1.38)
    check("s14", end + 0.28 + est_h(L["s14_note"], CW, 12.5, L["lang"], 1.38, 0, 0), BODY_BOT + 0.2)


def s15_conclusions(prs, L, n):
    s = base_slide(prs, L["s15_t"], L["s15_k"], n, FOOT)
    y = BODY_TOP + 0.02
    gap = 0.34
    cw = (CW - gap) / 2
    ch = BODY_BOT - y
    for i, (t, items, col, mark) in enumerate([
            (L["s15_can_t"], L["s15_can"], GREEN, "✓"),
            (L["s15_cant_t"], L["s15_cant"], AMBER, "✕")]):
        x = ML + i * (cw + gap)
        rect(s, x, y, cw, ch, fill=CARD if i == 0 else WHITE, line=None if i == 0 else LINE,
             lw=1.2)
        rect(s, x, y, cw, 0.42, fill=col, line=None)
        tf = textbox(s, x + 0.26, y + 0.10, cw - 0.52, 0.3)
        para(tf, t, size=13, color=WHITE, bold=True, first=True, space_after=0)
        yy = y + 0.68
        for it in items:
            hh = est_h(it, cw - 0.92, 12.5, L["lang"], 1.32, 0, 0)
            tf = textbox(s, x + 0.30, yy, 0.26, 0.3)
            para(tf, mark, size=12, color=col, bold=True, first=True, space_after=0)
            tf = textbox(s, x + 0.62, yy, cw - 0.92, hh + 0.1)
            para(tf, it, size=12.5, color=INK, first=True, space_after=0, line=1.32)
            yy += hh + 0.20
        check("s15-%d" % i, yy - y, ch)


def s16_progress(prs, L, n):
    s = base_slide(prs, L["s16_t"], L["s16_k"], n, FOOT)
    y = BODY_TOP + 0.02
    gap = 0.26
    ws = [CW * 0.355, CW * 0.245, CW * 0.40]
    ws = [w - gap * 2 / 3 for w in ws]
    ch = BODY_BOT - y
    xs = [ML, ML + ws[0] + gap, ML + ws[0] + ws[1] + 2 * gap]
    trio = [(L["s16_done_t"], L["s16_done"], GREEN),
            (L["s16_wip_t"], L["s16_wip"], ACCENT),
            (L["s16_todo_t"], L["s16_todo"], AMBER)]
    for i, (t, items, col) in enumerate(trio):
        x, cw = xs[i], ws[i]
        rect(s, x, y, cw, ch, fill=CARD if i != 1 else CARD2, line=None)
        rect(s, x, y, cw, 0.40, fill=col, line=None)
        tf = textbox(s, x + 0.22, y + 0.09, cw - 0.44, 0.3)
        para(tf, t, size=12.5, color=WHITE, bold=True, first=True, space_after=0)
        yy = y + 0.62
        for it in items:
            hh = est_h(it, cw - 0.66, 11.5, L["lang"], 1.28, 0, 0)
            rect(s, x + 0.24, yy + 0.085, 0.05, 0.05, fill=col, line=None, shape=MSO_SHAPE.OVAL)
            tf = textbox(s, x + 0.44, yy, cw - 0.66, hh + 0.1)
            para(tf, it, size=11.5, color=INK, first=True, space_after=0, line=1.28)
            yy += hh + 0.135
        check("s16-%d" % i, yy - y, ch)


def s17_roadmap(prs, L, n):
    s = base_slide(prs, L["s17_t"], L["s17_k"], n, FOOT)
    end = grid(s, L, ML, BODY_TOP + 0.06, CW, [0.62, 2.3, 5.3, 1.5],
               L["s17_rows"], size=12.5, badge=True, first_bold=False,
               colcolor={1: INK, 3: ACCENT_D})
    tf = textbox(s, ML, end + 0.26, CW, 0.6)
    para(tf, L["s17_note"], size=12.5, color=SLATE, first=True, space_after=0, line=1.38)
    check("s17", end + 0.26 + est_h(L["s17_note"], CW, 12.5, L["lang"], 1.38, 0, 0), BODY_BOT + 0.2)


def s18_risks(prs, L, n):
    s = base_slide(prs, L["s18_t"], L["s18_k"], n, FOOT)
    y = BODY_TOP + 0.0
    for i, (title, why, fix) in enumerate(L["s18_rows"]):
        h1 = est_h(title, CW * 0.26 - 0.3, 13, L["lang"], 1.22, 0, 0)
        h2 = est_h(why, CW * 0.34 - 0.3, 11.5, L["lang"], 1.3, 0, 0)
        h3 = est_h(fix, CW * 0.40 - 0.3, 11.5, L["lang"], 1.3, 0, 0)
        rh = max(h1, h2, h3) + 0.34
        if i % 2 == 0:
            rect(s, ML, y, CW, rh, fill=CARD, line=None)
        rect(s, ML, y, 0.045, rh, fill=AMBER if i < 3 else SLATE, line=None)
        tf = textbox(s, ML + 0.26, y + 0.17, CW * 0.26 - 0.3, rh)
        para(tf, title, size=13, color=INK, bold=True, first=True, space_after=0, line=1.22)
        tf = textbox(s, ML + CW * 0.26 + 0.06, y + 0.17, CW * 0.34 - 0.3, rh)
        para(tf, why, size=11.5, color=MUTED, first=True, space_after=0, line=1.3)
        tf = textbox(s, ML + CW * 0.60 + 0.06, y + 0.17, CW * 0.40 - 0.24, rh)
        para(tf, fix, size=11.5, color=ACCENT_D, first=True, space_after=0, line=1.3)
        y += rh + 0.09
    check("s18", y, BODY_BOT + 0.15)


def s19_contrib(prs, L, n):
    s = base_slide(prs, L["s19_t"], L["s19_k"], n, FOOT)
    y = BODY_TOP + 0.0
    gap = 0.24
    cw = (CW - 3 * gap) / 4
    ch = 1.90
    for i, (head, body) in enumerate(L["s19_cards"]):
        x = ML + i * (cw + gap)
        rect(s, x, y, cw, ch, fill=WHITE, line=ACCENT if i < 2 else AMBER, lw=1.4)
        tf = textbox(s, x + 0.22, y + 0.24, cw - 0.44, 0.6)
        para(tf, head, size=13.5, color=INK, bold=True, first=True, space_after=0, line=1.18)
        hh = est_h(head, cw - 0.44, 13.5, L["lang"], 1.18, 0, 0)
        tf = textbox(s, x + 0.22, y + 0.28 + hh + 0.12, cw - 0.44, ch - hh - 0.5)
        para(tf, body, size=11.5, color=SLATE, first=True, space_after=0, line=1.32)
        check("s19c-%d" % i, hh + est_h(body, cw - 0.44, 11.5, L["lang"], 1.32, 0, 0) + 0.5, ch)

    y2 = y + ch + 0.26
    gap2 = 0.34
    w1 = CW * 0.53 - gap2 / 2
    w2 = CW - w1 - gap2
    ch2 = BODY_BOT - y2
    for x, w, t, items, col in [(ML, w1, L["s19_out_t"], L["s19_out"], SLATE),
                                (ML + w1 + gap2, w2, L["s19_ask_t"], L["s19_ask"], ACCENT_D)]:
        rect(s, x, y2, w, ch2, fill=CARD if col == SLATE else ACCENT_D, line=None)
        fg = WHITE if col == ACCENT_D else INK
        tf = textbox(s, x + 0.26, y2 + 0.22, w - 0.52, 0.3)
        para(tf, t, size=12, color=(RGBColor(0x8E, 0xD6, 0xE4) if col == ACCENT_D else ACCENT_D),
             bold=True, first=True, space_after=0)
        yy = y2 + 0.66
        for it in items:
            hh = est_h(it, w - 0.86, 12, L["lang"], 1.3, 0, 0)
            rect(s, x + 0.28, yy + 0.085, 0.05, 0.05,
                 fill=(AMBER if col == ACCENT_D else ACCENT), line=None, shape=MSO_SHAPE.OVAL)
            tf = textbox(s, x + 0.50, yy, w - 0.8, hh + 0.1)
            para(tf, it, size=12, color=fg, first=True, space_after=0, line=1.3)
            yy += hh + 0.16
        check("s19-%s" % t[:6], yy - y2, ch2)


# ---------------------------------------------------------------- driver
def build(L, outpath):
    global FOOT
    FOOT = "HTN-Concord · " + ("博士课题设计与进度" if L["lang"] == "zh"
                               else "PhD project design and progress")
    prs = new_deck()
    cover(prs, L)
    fns = [s2_quad, s3_constraints, s4_gap, s5_positioning, s6_engine, s7_subtraction,
           s8_arch, s9_scope, s10_data, s11_tasks, s12_ladder, s13_metrics,
           s14_expected, s15_conclusions, s16_progress, s17_roadmap, s18_risks,
           s19_contrib]
    for i, fn in enumerate(fns, start=2):
        fn(prs, L, i)
    prs.save(outpath)
    return len(prs.slides.__iter__.__self__._sldIdLst)


if __name__ == "__main__":
    outdir = sys.argv[1]
    for L in (C.EN, C.ZH):
        WARN[:] = []
        p = os.path.join(outdir, L["file"])
        build(L, p)
        print("%-42s %2d slides  %6.1f KB" % (L["file"], 19, os.path.getsize(p) / 1024))
        for w in WARN:
            print("    OVERFLOW  " + w)
        if not WARN:
            print("    layout check: no overflow")
